"""Builds synthetic test documents in every supported format. All values are fake."""

from __future__ import annotations

import io
import json
import math
import zipfile
from pathlib import Path

import docx
import pymupdf
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont

from app.processing.detection.validators.verhoeff import verhoeff_check_digit


def aadhaar(base: str = "48217730551") -> str:
    return base + verhoeff_check_digit(base)


AADHAAR = aadhaar()
AADHAAR_SPACED = f"{AADHAAR[:4]} {AADHAAR[4:8]} {AADHAAR[8:]}"
PAN = "ABCPK1234F"
PHONE = "9845021773"
EMAIL = "rohan.kumar@example.in"
NAME = "Rohan Kumar"
PASSPORT = "N4821773"

FORM_LINES = [
    "Employee application form",
    f"Applicant name: {NAME}",
    f"Aadhaar No: {AADHAAR_SPACED}",
    f"PAN: {PAN}",
    f"Mobile: +91 {PHONE[:5]} {PHONE[5:]}",
    f"Email: {EMAIL}",
    "Invoice Number: 123456789012",
]
RAW_VALUES = [NAME, AADHAAR_SPACED, AADHAAR, PAN, PHONE, EMAIL]


def txt_bytes(lines: list[str] | None = None) -> bytes:
    return ("\n".join(lines or FORM_LINES) + "\n").encode()


def pdf_bytes(lines: list[str] | None = None, *, pages: int = 1) -> bytes:
    document = pymupdf.open()
    for _ in range(pages):
        page = document.new_page()
        y = 72
        for line in lines or FORM_LINES:
            page.insert_text((60, y), line, fontsize=11)
            y += 20
    data: bytes = document.tobytes()
    document.close()
    return data


def text_image(lines: list[str] | None = None) -> Image.Image:
    font = ImageFont.load_default(size=30)
    content = lines or FORM_LINES[:6]
    image = Image.new("RGB", (1700, 120 + 52 * len(content)), "white")
    draw = ImageDraw.Draw(image)
    for index, line in enumerate(content):
        draw.text((70, 60 + 52 * index), line, fill="black", font=font)
    return image


def png_bytes(lines: list[str] | None = None) -> bytes:
    buffer = io.BytesIO()
    text_image(lines).save(buffer, format="PNG")
    return buffer.getvalue()


def scanned_pdf_bytes(lines: list[str] | None = None) -> bytes:
    image = text_image(lines)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    document = pymupdf.open()
    page = document.new_page(width=612, height=612 * image.height / image.width)
    page.insert_image(page.rect, stream=buffer.getvalue())
    data: bytes = document.tobytes()
    document.close()
    return data


def scanner_app_pdf_bytes(image_lines: list[str], hidden_lines: list[str]) -> bytes:
    """A page photo plus an invisible text layer, as scanner apps (Adobe Scan, CamScanner) produce.

    The two say different things on purpose, so a test can tell which reading a value came from.
    """
    image = text_image(image_lines)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    document = pymupdf.open()
    page = document.new_page(width=612, height=612 * image.height / image.width)
    page.insert_image(page.rect, stream=buffer.getvalue())
    for index, line in enumerate(hidden_lines):
        page.insert_text((30, 30 + 14 * index), line, fontsize=9, render_mode=3)
    data: bytes = document.tobytes()
    document.close()
    return data


def qr_image(payload: str = "https://verify.example.in/000000000001", side: int = 360) -> Image.Image:
    import cv2

    matrix = cv2.QRCodeEncoder.create().encode(payload)
    return Image.fromarray(matrix).convert("L").resize((side, side), Image.Resampling.NEAREST)


def signature_image(width: int = 460, height: int = 140) -> Image.Image:
    """Handwriting-like strokes (a looping line and a flourish); not anyone's signature."""
    image = Image.new("L", (width, height), 255)
    draw = ImageDraw.Draw(image)
    points = [
        (width * (0.06 + 0.86 * step / 89), height * (0.5 - 0.3 * math.sin(step / 89 * 12 * math.pi)))
        for step in range(90)
    ]
    draw.line(points, fill=0, width=6, joint="curve")
    draw.line([(0.15 * width, 0.82 * height), (0.75 * width, 0.62 * height)], fill=0, width=6)
    return image


def _png(image: Image.Image) -> bytes:
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


#: Where visual_pdf_bytes places its QR code and signature (PDF points).
VISUAL_PDF_QR = (380.0, 560.0, 480.0, 660.0)
VISUAL_PDF_SIGNATURE = (80.0, 300.0, 218.0, 342.0)


def visual_pdf_bytes() -> bytes:
    """A digitally issued certificate: text, an embedded QR code, and a signature image over its caption."""
    document = pymupdf.open()
    page = document.new_page()
    y = 72
    for line in FORM_LINES[:3]:
        page.insert_text((60, y), line, fontsize=11)
        y += 20
    page.insert_image(pymupdf.Rect(*VISUAL_PDF_QR), stream=_png(qr_image()))
    page.insert_image(pymupdf.Rect(*VISUAL_PDF_SIGNATURE), stream=_png(signature_image()))
    page.insert_text((80, VISUAL_PDF_SIGNATURE[3] + 16), "Signature of the holder", fontsize=10)
    data: bytes = document.tobytes()
    document.close()
    return data


def visual_png_bytes() -> bytes:
    """An ID page photographed or scanned: text, a dense QR code, and a signature over its caption."""
    page = Image.new("L", (1700, 2200), 255)
    draw = ImageDraw.Draw(page)
    font = ImageFont.load_default(size=30)
    for index, line in enumerate(FORM_LINES[:3]):
        draw.text((70, 200 + 52 * index), line, fill=0, font=font)
    page.paste(qr_image("1" * 1200, 400), (1180, 1300))
    strokes = signature_image()
    page.paste(strokes, (120, 700), mask=Image.eval(strokes, lambda value: 255 - value))
    draw.text((120, 860), "Signature of the holder", fill=0, font=font)
    return _png(page.convert("RGB"))


def docx_bytes(lines: list[str] | None = None) -> bytes:
    document = docx.Document()
    for line in lines or FORM_LINES:
        paragraph = document.add_paragraph()
        # Split into several runs so redaction must handle values spanning runs.
        middle = len(line) // 2
        paragraph.add_run(line[:middle])
        paragraph.add_run(line[middle:]).bold = True
    document.core_properties.author = NAME
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Employees"
    sheet.append(["Name", "PAN", "Aadhaar", "Mobile", "Invoice Number"])
    sheet.append([NAME, PAN, int(AADHAAR), PHONE, "123456789012"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def csv_bytes() -> bytes:
    return f"Name,PAN,Aadhaar,Mobile\n{NAME},{PAN},{AADHAAR},{PHONE}\n".encode()


def json_bytes() -> bytes:
    return json.dumps(
        {"customers": [{"full_name": NAME, "pan": PAN, "mobile": PHONE, "aadhaar": int(AADHAAR)}]}
    ).encode()


def xml_bytes() -> bytes:
    return (
        f"<?xml version='1.0'?><employees><employee><name>{NAME}</name><pan>{PAN}</pan>"
        f"<email>{EMAIL}</email></employee></employees>"
    ).encode()


BILLION_LAUGHS = b"""<?xml version="1.0"?>
<!DOCTYPE lolz [<!ENTITY lol "lol"><!ENTITY lol2 "&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;">]>
<lolz>&lol2;</lolz>"""


def zip_bomb_docx() -> bytes:
    """A DOCX-shaped archive whose single member expands enormously."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", b"0" * (20 * 1024 * 1024))
    return buffer.getvalue()


def write(tmp: Path, name: str, data: bytes) -> Path:
    path = tmp / name
    path.write_bytes(data)
    return path
