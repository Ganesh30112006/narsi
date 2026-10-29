"""DOCX extraction in document order: headers, body paragraphs and tables, footers."""

from __future__ import annotations

from pathlib import Path
from zipfile import BadZipFile

import docx
from docx.document import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.core.enums import FileType
from app.core.exceptions import InvalidFileError
from app.processing.extractor.common import pages_from_text
from app.processing.limits import ProcessingLimits
from app.processing.types import ExtractedDocument


def iter_block_text(document: DocxDocument) -> list[str]:
    """Paragraph text and table rows (as `cell | cell`) in body order. Shared with DOCX redaction."""
    lines: list[str] = []
    for section in document.sections:
        for paragraph in section.header.paragraphs:
            if paragraph.text.strip():
                lines.append(paragraph.text)
    body = document.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            lines.append(Paragraph(child, document).text)
        elif tag == "tbl":
            table = Table(child, document)
            for row in table.rows:
                seen: list[str] = []
                for cell in row.cells:
                    text = cell.text.strip()
                    # Merged cells repeat; keep them once.
                    if not seen or seen[-1] != text:
                        seen.append(text)
                lines.append(" | ".join(seen))
    for section in document.sections:
        for paragraph in section.footer.paragraphs:
            if paragraph.text.strip():
                lines.append(paragraph.text)
    return lines


def extract_docx(path: Path, limits: ProcessingLimits) -> ExtractedDocument:
    try:
        document = docx.Document(str(path))
    except (PackageNotFoundError, BadZipFile, KeyError, ValueError) as exc:
        raise InvalidFileError("This Word document could not be opened.") from exc
    text = "\n".join(iter_block_text(document))
    core = document.core_properties
    result = pages_from_text(text, source_type=FileType.DOCX, limits=limits)
    if core.title:
        result.metadata["title"] = core.title[:200]
    return result
